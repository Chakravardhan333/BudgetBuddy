from django.test import TestCase
from django.contrib.auth.models import User
from django.urls import reverse
from rest_framework.test import APIClient
from rest_framework import status
from decimal import Decimal
from datetime import date
from transactions.models import Expense, Income, Budget, SavingsGoal

class ReportsExportTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username='report_user', email='report@budgetbuddy.test', password='Password123!'
        )
        login_res = self.client.post(reverse('login'), {'username': 'report_user', 'password': 'Password123!'})
        self.token = login_res.data['access']
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {self.token}')

        # Seed sample data for the user
        Income.objects.create(
            user=self.user, amount=Decimal('2000.00'), source='salary',
            date=date.today(), description='Monthly Salary'
        )
        Expense.objects.create(
            user=self.user, amount=Decimal('150.00'), category_name='Food & Dining',
            date=date.today(), description='Dinner with family', payment_method='card'
        )
        Budget.objects.create(
            user=self.user, month=date.today().month, year=date.today().year,
            monthly_limit=Decimal('500.00')
        )
        SavingsGoal.objects.create(
            user=self.user, title='Laptop Fund', target_amount=Decimal('1000.00'),
            current_amount=Decimal('500.00'), target_date=date(2026, 12, 31)
        )

        self.pdf_url = reverse('export_pdf')
        self.excel_url = reverse('export_excel')

    def test_export_pdf_success(self):
        response = self.client.get(self.pdf_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-'))
        self.assertIn('attachment; filename=', response['Content-Disposition'])

    def test_export_excel_success(self):
        response = self.client.get(self.excel_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        # Check ZIP / XLSX signature
        self.assertTrue(response.content.startswith(b'PK\x03\x04'))
        self.assertIn('attachment; filename=', response['Content-Disposition'])

    def test_reports_require_authentication(self):
        unauth_client = APIClient()
        pdf_res = unauth_client.get(self.pdf_url)
        self.assertEqual(pdf_res.status_code, status.HTTP_401_UNAUTHORIZED)
        excel_res = unauth_client.get(self.excel_url)
        self.assertEqual(excel_res.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_export_reports_with_date_filters(self):
        today = date.today()
        res_pdf = self.client.get(self.pdf_url, {'year': today.year, 'month': today.month})
        self.assertEqual(res_pdf.status_code, status.HTTP_200_OK)
        self.assertTrue(res_pdf.content.startswith(b'%PDF-'))

        res_excel = self.client.get(self.excel_url, {'year': today.year, 'month': today.month})
        self.assertEqual(res_excel.status_code, status.HTTP_200_OK)
        self.assertTrue(res_excel.content.startswith(b'PK\x03\x04'))

    def test_user_data_isolation_in_reports(self):
        # Create second user with distinct secret expense
        other_user = User.objects.create_user(username='other_rep_user', password='Password123!')
        Expense.objects.create(
            user=other_user, amount=Decimal('9999.00'), category_name='SecretCategory',
            date=date.today(), description='ConfidentialOtherExpense'
        )

        # First user requests excel report
        excel_res = self.client.get(self.excel_url)
        import openpyxl, io
        wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
        # Ensure 'Budget Summary' sheet exists
        self.assertIn('Budget Summary', wb.sheetnames)
        self.assertIn('Expenses', wb.sheetnames)
        exp_sheet = wb['Expenses']
        descriptions = [row[2] for row in exp_sheet.iter_rows(min_row=2, values_only=True)]
        self.assertNotIn('ConfidentialOtherExpense', descriptions)
        self.assertIn('Dinner with family', descriptions)
