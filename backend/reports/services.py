import io
from datetime import datetime, date
from decimal import Decimal
from zoneinfo import ZoneInfo
from django.db.models import Sum

# ReportLab Imports
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# OpenPyXL Imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from transactions.models import Expense, Income, Budget, SavingsGoal

def get_user_currency_symbol(user):
    if hasattr(user, 'profile') and user.profile.currency_symbol:
        return user.profile.currency_symbol
    return '$'

def _get_localized_now_str(tz_name='Asia/Kolkata'):
    try:
        tz = ZoneInfo(tz_name) if tz_name else ZoneInfo('Asia/Kolkata')
    except Exception:
        tz = ZoneInfo('Asia/Kolkata')
    now = datetime.now(tz)
    return now.strftime("%B %d, %Y at %I:%M %p %Z")

def generate_pdf_report(user, year=None, month=None, timezone_str='Asia/Kolkata'):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    currency_sym = get_user_currency_symbol(user)
    currency_code = getattr(user.profile, 'currency', 'USD') if hasattr(user, 'profile') else 'USD'
    now_str = _get_localized_now_str(timezone_str)

    # Filter records
    expenses_qs = Expense.objects.filter(user=user)
    incomes_qs = Income.objects.filter(user=user)
    budgets_qs = Budget.objects.filter(user=user)
    goals_qs = SavingsGoal.objects.filter(user=user)

    period_title = "Complete Financial Statement"
    if year and month:
        expenses_qs = expenses_qs.filter(date__year=year, date__month=month)
        incomes_qs = incomes_qs.filter(date__year=year, date__month=month)
        budgets_qs = budgets_qs.filter(year=year, month=month)
        period_title = f"Statement for {date(year, month, 1).strftime('%B %Y')}"
    elif year:
        expenses_qs = expenses_qs.filter(date__year=year)
        incomes_qs = incomes_qs.filter(date__year=year)
        budgets_qs = budgets_qs.filter(year=year)
        period_title = f"Annual Statement for {year}"

    total_income = incomes_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_expense = expenses_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    net_savings = total_income - total_expense

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748b'),
    )

    section_heading = ParagraphStyle(
        'SectionHead',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0f172a'),
        spaceBefore=12,
        spaceAfter=6,
    )

    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1e293b'),
    )

    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0f172a'),
    )

    elements = []

    # Document Header
    header_data = [
        [
            Paragraph("<b>BudgetBuddy</b> &bull; Personal Finance Platform", title_style),
            Paragraph(f"<b>Generated:</b> {now_str}<br/><b>User:</b> {user.get_full_name() or user.username} ({user.email})", subtitle_style),
        ]
    ]
    t_header = Table(header_data, colWidths=[320, 220])
    t_header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
    ]))
    elements.append(t_header)
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#10b981'), spaceAfter=12))

    # Period Banner
    elements.append(Paragraph(f"<b>{period_title}</b> ({currency_code})", section_heading))

    # Executive Summary Table
    summary_data = [
        [Paragraph("Metric", cell_bold), Paragraph("Amount", cell_bold), Paragraph("Financial Implication", cell_bold)],
        [Paragraph("Total Cash Inflow", cell_style), Paragraph(f"{currency_sym}{total_income:.2f}", cell_bold), Paragraph("Gross income and funding sources", cell_style)],
        [Paragraph("Total Cash Outflow", cell_style), Paragraph(f"{currency_sym}{total_expense:.2f}", cell_bold), Paragraph("All recorded expenditures", cell_style)],
        [Paragraph("Net Savings / Surplus", cell_style), Paragraph(f"{currency_sym}{net_savings:.2f}", cell_bold), Paragraph("Retained financial liquidity", cell_style)],
    ]
    t_summary = Table(summary_data, colWidths=[180, 140, 220])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
    ]))
    elements.append(t_summary)
    elements.append(Spacer(1, 12))

    # Category Spending Breakdown
    cat_qs = expenses_qs.values('category_name').annotate(total=Sum('amount')).order_by('-total')
    if cat_qs.exists():
        elements.append(Paragraph("Category-Wise Expense Distribution", section_heading))
        cat_data = [[Paragraph("Category", cell_bold), Paragraph("Amount", cell_bold), Paragraph("Share of Expenses", cell_bold)]]
        for c in cat_qs:
            share = (c['total'] / total_expense * 100) if total_expense > 0 else 0
            cat_data.append([
                Paragraph(c['category_name'], cell_style),
                Paragraph(f"{currency_sym}{c['total']:.2f}", cell_style),
                Paragraph(f"{share:.1f}%", cell_style),
            ])
        t_cat = Table(cat_data, colWidths=[240, 150, 150])
        t_cat.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_cat)
        elements.append(Spacer(1, 12))

    # Budget Utilization Summary
    if budgets_qs.exists():
        elements.append(Paragraph("Monthly Budget Utilization & Performance", section_heading))
        b_data = [[Paragraph("Period", cell_bold), Paragraph("Scope / Category", cell_bold), Paragraph("Budget Cap", cell_bold), Paragraph("Actual Spending", cell_bold), Paragraph("Utilization", cell_bold)]]
        for b in budgets_qs:
            exp_b = Expense.objects.filter(user=user, date__year=b.year, date__month=b.month)
            if b.category:
                exp_b = exp_b.filter(category_name__iexact=b.category)
            spending = exp_b.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
            util = float(spending / b.monthly_limit * 100) if b.monthly_limit > 0 else 0.0
            b_data.append([
                Paragraph(f"{b.month:02d}/{b.year}", cell_style),
                Paragraph(b.category or "Overall Monthly", cell_style),
                Paragraph(f"{currency_sym}{b.monthly_limit:.2f}", cell_style),
                Paragraph(f"{currency_sym}{spending:.2f}", cell_style),
                Paragraph(f"{util:.1f}%", cell_bold if util >= b.warning_threshold else cell_style),
            ])
        t_b = Table(b_data, colWidths=[80, 180, 100, 100, 80])
        t_b.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_b)
        elements.append(Spacer(1, 12))

    # Savings Goals Summary
    if goals_qs.exists():
        elements.append(Paragraph("Savings Goals Progress", section_heading))
        goals_data = [[Paragraph("Goal Title", cell_bold), Paragraph("Target Amount", cell_bold), Paragraph("Saved Amount", cell_bold), Paragraph("Progress", cell_bold), Paragraph("Status", cell_bold)]]
        for g in goals_qs:
            pct = (g.current_amount / g.target_amount * 100) if g.target_amount > 0 else 0
            goals_data.append([
                Paragraph(g.title, cell_style),
                Paragraph(f"{currency_sym}{g.target_amount:.2f}", cell_style),
                Paragraph(f"{currency_sym}{g.current_amount:.2f}", cell_style),
                Paragraph(f"{pct:.1f}%", cell_style),
                Paragraph(g.get_status_display(), cell_style),
            ])
        t_goals = Table(goals_data, colWidths=[180, 90, 90, 80, 100])
        t_goals.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f8fafc')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(t_goals)
        elements.append(Spacer(1, 12))

    # Recent Transactions Ledger
    elements.append(Paragraph("Recent Transaction Ledger (Debits & Credits)", section_heading))
    combined = []
    for e in expenses_qs[:15]:
        combined.append({
            'date': e.date,
            'desc': e.description,
            'cat': e.category_name,
            'type': 'Expense',
            'amount': f"-{currency_sym}{e.amount:.2f}",
        })
    for i in incomes_qs[:15]:
        combined.append({
            'date': i.date,
            'desc': i.description,
            'cat': i.get_source_display(),
            'type': 'Income',
            'amount': f"+{currency_sym}{i.amount:.2f}",
        })
    combined.sort(key=lambda x: x['date'], reverse=True)

    tx_data = [[Paragraph("Date", cell_bold), Paragraph("Description", cell_bold), Paragraph("Category / Source", cell_bold), Paragraph("Type", cell_bold), Paragraph("Amount", cell_bold)]]
    for row in combined[:20]:
        tx_data.append([
            Paragraph(row['date'].strftime("%Y-%m-%d"), cell_style),
            Paragraph(row['desc'], cell_style),
            Paragraph(row['cat'], cell_style),
            Paragraph(row['type'], cell_style),
            Paragraph(row['amount'], cell_bold),
        ])

    t_tx = Table(tx_data, colWidths=[70, 180, 130, 70, 90])
    t_tx.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_tx)

    # Footer note
    elements.append(Spacer(1, 18))
    elements.append(Paragraph("Generated securely by BudgetBuddy Personal Finance Platform &bull; Data is confidential and isolated.", subtitle_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(user, year=None, month=None, timezone_str='Asia/Kolkata'):
    wb = openpyxl.Workbook()
    currency_sym = get_user_currency_symbol(user)

    # Styling definitions
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0'),
    )

    # Filter querysets
    expenses_qs = Expense.objects.filter(user=user)
    incomes_qs = Income.objects.filter(user=user)
    budgets_qs = Budget.objects.filter(user=user)
    goals_qs = SavingsGoal.objects.filter(user=user)

    if year and month:
        expenses_qs = expenses_qs.filter(date__year=year, date__month=month)
        incomes_qs = incomes_qs.filter(date__year=year, date__month=month)
        budgets_qs = budgets_qs.filter(year=year, month=month)
    elif year:
        expenses_qs = expenses_qs.filter(date__year=year)
        incomes_qs = incomes_qs.filter(date__year=year)
        budgets_qs = budgets_qs.filter(year=year)

    total_income = incomes_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    total_expense = expenses_qs.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
    net_savings = total_income - total_expense

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.views.sheetView[0].showGridLines = True

    ws_summary["A1"] = "BudgetBuddy Financial Statement"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Account: {user.get_full_name() or user.username} ({user.email})"
    ws_summary["A3"] = f"Generated: {_get_localized_now_str(timezone_str)}"
    ws_summary["A4"] = f"Currency: {currency_sym} ({getattr(user.profile, 'currency', 'USD') if hasattr(user, 'profile') else 'USD'})"

    summary_rows = [
        ["Key Metric", "Total Value", "Notes"],
        ["Total Incomes", float(total_income), "All income streams logged"],
        ["Total Expenses", float(total_expense), "All expenditures logged"],
        ["Net Cash Flow / Savings", float(net_savings), "Retained financial surplus"],
    ]

    for r_idx, row in enumerate(summary_rows, start=6):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            if r_idx == 6:
                cell.fill = navy_fill
                cell.font = header_font
            else:
                cell.font = bold_font if c_idx == 2 else regular_font

    # Sheet 2: Expenses
    ws_expenses = wb.create_sheet(title="Expenses")
    ws_expenses.views.sheetView[0].showGridLines = True
    exp_headers = ["ID", "Date", "Description", "Category", "Payment Method", "Amount", "Notes"]
    ws_expenses.append(exp_headers)
    for c_idx in range(1, len(exp_headers) + 1):
        cell = ws_expenses.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    for e in expenses_qs:
        ws_expenses.append([
            e.id,
            e.date.strftime("%Y-%m-%d"),
            e.description,
            e.category_name,
            e.get_payment_method_display(),
            float(e.amount),
            e.notes or "",
        ])

    # Sheet 3: Income
    ws_income = wb.create_sheet(title="Income")
    ws_income.views.sheetView[0].showGridLines = True
    inc_headers = ["ID", "Date", "Description", "Source", "Amount", "Notes"]
    ws_income.append(inc_headers)
    for c_idx in range(1, len(inc_headers) + 1):
        cell = ws_income.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    for i in incomes_qs:
        ws_income.append([
            i.id,
            i.date.strftime("%Y-%m-%d"),
            i.description,
            i.get_source_display(),
            float(i.amount),
            i.notes or "",
        ])

    # Sheet 4: Categories
    ws_categories = wb.create_sheet(title="Categories")
    ws_categories.views.sheetView[0].showGridLines = True
    cat_headers = ["Category Name", "Total Outlay", "Percentage Share"]
    ws_categories.append(cat_headers)
    for c_idx in range(1, len(cat_headers) + 1):
        cell = ws_categories.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    cat_qs = expenses_qs.values('category_name').annotate(total=Sum('amount')).order_by('-total')
    for c in cat_qs:
        share = float(c['total'] / total_expense * 100) if total_expense > 0 else 0.0
        ws_categories.append([c['category_name'], float(c['total']), f"{share:.2f}%"])

    # Sheet 5: Savings Goals
    ws_goals = wb.create_sheet(title="Savings Goals")
    ws_goals.views.sheetView[0].showGridLines = True
    goal_headers = ["ID", "Goal Title", "Target Amount", "Current Saved", "Remaining", "Deadline", "Status", "Progress %"]
    ws_goals.append(goal_headers)
    for c_idx in range(1, len(goal_headers) + 1):
        cell = ws_goals.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    for g in goals_qs:
        pct = float(g.current_amount / g.target_amount * 100) if g.target_amount > 0 else 0.0
        ws_goals.append([
            g.id,
            g.title,
            float(g.target_amount),
            float(g.current_amount),
            float(max(0, g.target_amount - g.current_amount)),
            g.target_date.strftime("%Y-%m-%d"),
            g.get_status_display(),
            f"{pct:.1f}%",
        ])

    # Sheet 6: Budget Summary
    ws_budgets = wb.create_sheet(title="Budget Summary")
    ws_budgets.views.sheetView[0].showGridLines = True
    budget_headers = ["Period", "Category", "Budget Limit", "Actual Spending", "Remaining", "Utilization %", "Status"]
    ws_budgets.append(budget_headers)
    for c_idx in range(1, len(budget_headers) + 1):
        cell = ws_budgets.cell(row=1, column=c_idx)
        cell.fill = navy_fill
        cell.font = header_font

    for b in budgets_qs:
        exp_b = Expense.objects.filter(user=user, date__year=b.year, date__month=b.month)
        if b.category:
            exp_b = exp_b.filter(category_name__iexact=b.category)
        spending = exp_b.aggregate(total=Sum('amount'))['total'] or Decimal('0.00')
        remaining = b.monthly_limit - spending
        util = float(spending / b.monthly_limit * 100) if b.monthly_limit > 0 else 0.0
        status_label = "Exceeded" if spending > b.monthly_limit else ("Near Limit" if util >= b.warning_threshold else "On Track")
        ws_budgets.append([
            f"{b.month:02d}/{b.year}",
            b.category or "Overall Monthly",
            float(b.monthly_limit),
            float(spending),
            float(remaining),
            f"{util:.1f}%",
            status_label,
        ])

    # Auto-adjust column widths for all sheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
